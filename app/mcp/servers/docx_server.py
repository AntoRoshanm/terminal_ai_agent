"""
DOCX Document Generation MCP Server (Stdio JSON-RPC 2.0)
Wraps python-docx for DOCX document creation with B.40 verification.
"""

import json
from pathlib import Path
import sys
from typing import Any, Dict, List, Optional

# Ensure UTF-8 I/O on Windows when available
if sys.platform == "win32":
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8")
            except Exception:
                pass


def log_debug(msg: str) -> None:
    try:
        sys.stderr.write(f"[MCP DOCX Server] {msg}\n")
        sys.stderr.flush()
    except Exception:
        pass


def generate_docx(file_path: str, title: str, content: str = "", sections: Optional[List[Dict[str, str]]] = None) -> Dict[str, Any]:
    """Generate a formatted DOCX document using python-docx."""
    path = Path(file_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor

        doc = Document()

        # Add Title
        title_para = doc.add_heading(title, level=0)
        title_para.runs[0].font.size = Pt(22)
        title_para.runs[0].font.color.rgb = RGBColor(26, 54, 93)

        # Explicit sections if provided
        if sections:
            for sec in sections:
                h = sec.get("heading", sec.get("title", ""))
                b = sec.get("body", sec.get("content", ""))
                if h:
                    doc.add_heading(h, level=1)
                if b:
                    for para in b.split("\n\n"):
                        if para.strip():
                            doc.add_paragraph(para.strip())

        # Parse markdown content
        if content:
            for line in content.split("\n"):
                clean = line.strip()
                if not clean:
                    continue
                elif clean.startswith("### "):
                    doc.add_heading(clean[4:], level=3)
                elif clean.startswith("## "):
                    doc.add_heading(clean[3:], level=2)
                elif clean.startswith("# "):
                    doc.add_heading(clean[2:], level=1)
                elif clean.startswith("- ") or clean.startswith("* "):
                    doc.add_paragraph(clean[2:], style="List Bullet")
                else:
                    doc.add_paragraph(clean)

        doc.save(str(path))

    except Exception as e:
        log_debug(f"python-docx generation error: {e}")
        raise RuntimeError(f"Failed to generate DOCX document: {e}")

    # Post-action verification (B.40)
    if not path.exists():
        raise FileNotFoundError(f"Verification Failed: DOCX file '{path}' was not written.")

    size_bytes = path.stat().st_size
    if size_bytes == 0:
        raise IOError(f"Verification Failed: DOCX file '{path}' is 0 bytes.")

    return {
        "file_path": str(path),
        "size_bytes": size_bytes,
        "title": title,
        "status": "success",
        "verified_success": True,
    }


def handle_request(req: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    method = req.get("method")
    req_id = req.get("id")

    if method == "initialize":
        return {
            "jsonrpc": "2.0",
            "id": req_id,
            "result": {
                "protocolVersion": "2024-11-05",
                "capabilities": {"tools": {}},
                "serverInfo": {
                    "name": "docx_document_server",
                    "version": "1.0.0",
                },
            },
        }

    elif method == "notifications/initialized":
        return None

    elif method == "tools/list":
        return {
            "jsonrpc": "2.0",
            "id": req_id,
            "result": {
                "tools": [
                    {
                        "name": "create_docx",
                        "description": "Create a formatted, verified DOCX Microsoft Word document with headings and styled paragraphs.",
                        "inputSchema": {
                            "type": "object",
                            "properties": {
                                "file_path": {
                                    "type": "string",
                                    "description": "Absolute destination path for the generated DOCX file",
                                },
                                "title": {
                                    "type": "string",
                                    "description": "Document title heading",
                                },
                                "content": {
                                    "type": "string",
                                    "description": "Complete text or markdown content of the document",
                                },
                                "sections": {
                                    "type": "array",
                                    "items": {
                                        "type": "object",
                                        "properties": {
                                            "heading": {"type": "string"},
                                            "body": {"type": "string"},
                                        },
                                    },
                                    "description": "Optional list of structured sections with heading and body",
                                },
                            },
                            "required": ["file_path", "title"],
                        },
                    }
                ]
            },
        }

    elif method == "tools/call":
        params = req.get("params", {})
        tool_name = params.get("name")
        arguments = params.get("arguments", {})

        if tool_name == "create_docx":
            file_path = arguments.get("file_path", "")
            title = arguments.get("title", "Document")
            content = arguments.get("content", "")
            sections = arguments.get("sections")

            try:
                res = generate_docx(file_path, title, content=content, sections=sections)
                return {
                    "jsonrpc": "2.0",
                    "id": req_id,
                    "result": {
                        "content": [
                            {
                                "type": "text",
                                "text": json.dumps(res, indent=2),
                            }
                        ],
                        "isError": False,
                    },
                }
            except Exception as e:
                return {
                    "jsonrpc": "2.0",
                    "id": req_id,
                    "result": {
                        "content": [{"type": "text", "text": f"Error generating DOCX: {e}"}],
                        "isError": True,
                    },
                }

        else:
            return {
                "jsonrpc": "2.0",
                "id": req_id,
                "error": {
                    "code": -32601,
                    "message": f"Unknown tool '{tool_name}'",
                },
            }

    else:
        return {
            "jsonrpc": "2.0",
            "id": req_id,
            "error": {
                "code": -32601,
                "message": f"Method '{method}' not found",
            },
        }


def main():
    log_debug("DOCX Document MCP Server started over stdio.")
    for line in sys.stdin:
        line = line.strip()
        if not line:
            continue
        try:
            req = json.loads(line)
            resp = handle_request(req)
            if resp is not None:
                sys.stdout.write(json.dumps(resp) + "\n")
                sys.stdout.flush()
        except Exception as e:
            log_debug(f"Error handling request: {e}")
            err_resp = {
                "jsonrpc": "2.0",
                "id": None,
                "error": {"code": -32700, "message": f"Parse error: {e}"},
            }
            sys.stdout.write(json.dumps(err_resp) + "\n")
            sys.stdout.flush()


if __name__ == "__main__":
    main()
